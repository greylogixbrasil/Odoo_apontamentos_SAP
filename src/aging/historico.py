"""Importacao unica da base historica do SharePoint (oportunidades ate 2025) - export CSV ou
Excel da lista. Idempotente: upsert por (base_origem='SharePoint', ID da lista)."""
import csv
import logging
from pathlib import Path

from openpyxl import load_workbook

from src.aging.mapeamento import mapear_sharepoint
from src.aging.repositorio import AgingRepositorio

logger = logging.getLogger(__name__)

COLUNAS_OBRIGATORIAS = {"ID", "Código", "Título", "Cliente", "Responsavel", "Vertical", "Estagio", "Criado",
                        "Modificado", "FechamentoReal", "ValEstimado", "ValFechado"}


def ler_export(caminho: str | Path) -> list[dict]:
    caminho = Path(caminho)
    if caminho.suffix.lower() in (".xlsx", ".xlsm"):
        planilha = load_workbook(caminho, read_only=True, data_only=True).active
        linhas = planilha.iter_rows(values_only=True)
        cabecalho = [str(c).strip() if c is not None else "" for c in next(linhas)]
        registros = [dict(zip(cabecalho, linha)) for linha in linhas if any(v is not None for v in linha)]
    else:
        with open(caminho, encoding="utf-8-sig", newline="") as fh:
            amostra = fh.read(4096)
            fh.seek(0)
            dialeto = csv.Sniffer().sniff(amostra, delimiters=",;\t")
            registros = list(csv.DictReader(fh, dialect=dialeto))

    faltando = COLUNAS_OBRIGATORIAS - set(registros[0].keys() if registros else [])
    if faltando:
        raise ValueError(f"export do SharePoint sem as colunas: {', '.join(sorted(faltando))}")
    return registros


def importar_historico(dsn: str, caminho: str | Path) -> dict:
    registros = [mapear_sharepoint(linha) for linha in ler_export(caminho)]
    repo = AgingRepositorio(dsn)
    log_id = repo.iniciar_log("historico")
    try:
        inseridos, atualizados = repo.upsert(registros)
        repo.finalizar_log(log_id, "ok", len(registros), inseridos, atualizados)
        resumo = {"lidos": len(registros), "inseridos": inseridos, "atualizados": atualizados}
        logger.info("importacao historica do SharePoint concluida: %s", resumo)
        return resumo
    except Exception as exc:
        repo.finalizar_log(log_id, "falhou", erro=str(exc)[:2000])
        raise
    finally:
        repo.close()
